import os
import io
import calendar
from datetime import datetime
from dotenv import load_dotenv
from flask import Flask, request, jsonify, send_from_directory, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from models import db, InventoryItem, Receivable, Payment, ProfitLog, CATEGORIES, CATEGORY_LABELS

load_dotenv()


def r2(value):
    """Round a number to 2 decimal places, safely handling None/blank input."""
    try:
        return round(float(value or 0), 2)
    except (TypeError, ValueError):
        return 0.0


BASE_DIR = os.path.abspath(os.path.dirname(__file__))
FRONTEND_DIR = os.path.join(BASE_DIR, "static")

app = Flask(__name__, static_folder=None)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "change-this-secret-key")

database_url = os.environ.get("DATABASE_URL")
if database_url:
    # Neon/most Postgres hosts give a "postgres://" or "postgresql://" URL;
    # SQLAlchemy wants the "postgresql://" form.
    if database_url.startswith("postgres://"):
        database_url = database_url.replace("postgres://", "postgresql://", 1)
    app.config["SQLALCHEMY_DATABASE_URI"] = database_url
else:
    # No DATABASE_URL set (e.g. plain local development) -> fall back to
    # SQLite. Vercel's filesystem is read-only except /tmp, so use that
    # there; use a normal local file everywhere else. Once DATABASE_URL
    # (Postgres) is set, this branch is never used on Vercel.
    if os.environ.get("VERCEL"):
        db_path = "/tmp/vendor.db"
    else:
        db_path = os.path.join(BASE_DIR, "vendor.db")
    app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///" + db_path

app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

db.init_app(app)



# ---------------------------------------------------------------------------
# Frontend page serving (static HTML/CSS/JS) - no login required
# ---------------------------------------------------------------------------

@app.route("/")
def serve_root():
    return send_from_directory(FRONTEND_DIR, "dashboard.html")


@app.route("/<path:filename>")
def serve_static_files(filename):
    full_path = os.path.join(FRONTEND_DIR, filename)
    if os.path.isfile(full_path):
        return send_from_directory(FRONTEND_DIR, filename)
    return jsonify({"error": "not_found"}), 404


# ---------------------------------------------------------------------------
# Inventory API
# ---------------------------------------------------------------------------

def log_profit_snapshot(item):
    """Record a profit snapshot for this item. Called whenever an item is
    created or updated, so you get a history of margin changes over time."""
    snapshot = ProfitLog(
        item_id=item.id,
        item_name=item.name,
        category=item.category,
        quantity=item.quantity,
        purchase_price=item.purchase_price,
        price_per_unit=item.price_per_unit,
        profit_per_unit=item.profit_per_unit,
        total_potential_profit=item.total_potential_profit,
    )
    db.session.add(snapshot)


@app.route("/api/inventory", methods=["GET", "POST"])
def api_inventory():
    if request.method == "POST":
        data = request.get_json(force=True)
        category = data.get("category")
        if category not in CATEGORIES:
            return jsonify({"ok": False, "error": "invalid category"}), 400
        item = InventoryItem(
            category=category,
            name=(data.get("name") or "").strip(),
            unit=data.get("unit") or "pcs",
            quantity=r2(data.get("quantity")),
            purchase_price=r2(data.get("purchase_price")),
            price_per_unit=r2(data.get("price_per_unit")),
            low_stock_threshold=r2(data.get("low_stock_threshold")),
            notes=data.get("notes") or "",
        )
        if not item.name:
            return jsonify({"ok": False, "error": "name is required"}), 400
        db.session.add(item)
        db.session.flush()  # assigns item.id before we log the first snapshot
        log_profit_snapshot(item)
        db.session.commit()
        return jsonify({"ok": True, "item": item.to_dict()})

    category = request.args.get("category")
    q = InventoryItem.query.filter_by(is_deleted=False)
    if category and category in CATEGORIES:
        q = q.filter_by(category=category)
    search = request.args.get("search")
    if search:
        q = q.filter(InventoryItem.name.ilike(f"%{search}%"))
    items = q.order_by(InventoryItem.category, InventoryItem.name).all()
    return jsonify([i.to_dict() for i in items])


@app.route("/api/inventory/trash")
def api_inventory_trash():
    items = (
        InventoryItem.query.filter_by(is_deleted=True)
        .order_by(InventoryItem.deleted_at.desc())
        .all()
    )
    return jsonify([i.to_dict() for i in items])


@app.route("/api/inventory/<int:item_id>/restore", methods=["POST"])
def api_inventory_restore(item_id):
    item = db.session.get(InventoryItem, item_id)
    if not item:
        return jsonify({"ok": False, "error": "not found"}), 404
    item.is_deleted = False
    item.deleted_at = None
    db.session.commit()
    return jsonify({"ok": True, "item": item.to_dict()})


@app.route("/api/inventory/<int:item_id>/permanent", methods=["DELETE"])
def api_inventory_permanent_delete(item_id):
    item = db.session.get(InventoryItem, item_id)
    if not item:
        return jsonify({"ok": False, "error": "not found"}), 404
    db.session.delete(item)
    db.session.commit()
    return jsonify({"ok": True})


@app.route("/api/inventory/<int:item_id>", methods=["PUT", "DELETE"])
def api_inventory_item(item_id):
    item = db.session.get(InventoryItem, item_id)
    if not item:
        return jsonify({"ok": False, "error": "not found"}), 404
    if request.method == "DELETE":
        item.is_deleted = True
        item.deleted_at = datetime.utcnow()
        db.session.commit()
        return jsonify({"ok": True})

    data = request.get_json(force=True)
    for field in ["name", "unit", "notes"]:
        if field in data:
            setattr(item, field, data[field])
    for field in ["quantity", "purchase_price", "price_per_unit", "low_stock_threshold"]:
        if field in data:
            setattr(item, field, r2(data[field]))
    if data.get("category") in CATEGORIES:
        item.category = data["category"]
    item.updated_at = datetime.utcnow()
    log_profit_snapshot(item)
    db.session.commit()
    return jsonify({"ok": True, "item": item.to_dict()})


@app.route("/api/inventory/<int:item_id>/profit-history")
def api_item_profit_history(item_id):
    logs = (
        ProfitLog.query.filter_by(item_id=item_id)
        .order_by(ProfitLog.recorded_at.desc())
        .all()
    )
    return jsonify([l.to_dict() for l in logs])


# ---------------------------------------------------------------------------
# Receivables API (money customers owe the vendor)
# ---------------------------------------------------------------------------

@app.route("/api/receivables", methods=["GET", "POST"])
def api_receivables():
    if request.method == "POST":
        data = request.get_json(force=True)
        name = (data.get("customer_name") or "").strip()
        if not name:
            return jsonify({"ok": False, "error": "customer_name is required"}), 400
        due_date = None
        if data.get("due_date"):
            due_date = datetime.strptime(data["due_date"], "%Y-%m-%d").date()
        r = Receivable(
            customer_name=name,
            phone=data.get("phone") or "",
            total_amount=r2(data.get("total_amount")),
            amount_paid=r2(data.get("amount_paid")),
            due_date=due_date,
            notes=data.get("notes") or "",
        )
        db.session.add(r)
        db.session.commit()
        return jsonify({"ok": True, "receivable": r.to_dict()})

    status_filter = request.args.get("status")
    search = request.args.get("search")
    q = Receivable.query.filter_by(is_deleted=False)
    if search:
        q = q.filter(Receivable.customer_name.ilike(f"%{search}%"))
    items = [r.to_dict() for r in q.order_by(Receivable.date_created.desc()).all()]
    if status_filter in ("pending", "partial", "paid"):
        items = [i for i in items if i["status"] == status_filter]
    return jsonify(items)


@app.route("/api/receivables/trash")
def api_receivables_trash():
    items = (
        Receivable.query.filter_by(is_deleted=True)
        .order_by(Receivable.deleted_at.desc())
        .all()
    )
    return jsonify([r.to_dict() for r in items])


@app.route("/api/receivables/<int:rec_id>/restore", methods=["POST"])
def api_receivable_restore(rec_id):
    r = db.session.get(Receivable, rec_id)
    if not r:
        return jsonify({"ok": False, "error": "not found"}), 404
    r.is_deleted = False
    r.deleted_at = None
    db.session.commit()
    return jsonify({"ok": True, "receivable": r.to_dict()})


@app.route("/api/receivables/<int:rec_id>/permanent", methods=["DELETE"])
def api_receivable_permanent_delete(rec_id):
    r = db.session.get(Receivable, rec_id)
    if not r:
        return jsonify({"ok": False, "error": "not found"}), 404
    db.session.delete(r)
    db.session.commit()
    return jsonify({"ok": True})


@app.route("/api/receivables/<int:rec_id>", methods=["PUT", "DELETE"])
def api_receivable_item(rec_id):
    r = db.session.get(Receivable, rec_id)
    if not r:
        return jsonify({"ok": False, "error": "not found"}), 404
    if request.method == "DELETE":
        r.is_deleted = True
        r.deleted_at = datetime.utcnow()
        db.session.commit()
        return jsonify({"ok": True})

    data = request.get_json(force=True)
    for field in ["customer_name", "phone", "notes"]:
        if field in data:
            setattr(r, field, data[field])
    if "total_amount" in data:
        r.total_amount = r2(data["total_amount"])
    if "due_date" in data:
        r.due_date = datetime.strptime(data["due_date"], "%Y-%m-%d").date() if data["due_date"] else None
    db.session.commit()
    return jsonify({"ok": True, "receivable": r.to_dict()})


@app.route("/api/receivables/<int:rec_id>/payment", methods=["POST"])
def api_receivable_payment(rec_id):
    r = db.session.get(Receivable, rec_id)
    if not r:
        return jsonify({"ok": False, "error": "not found"}), 404
    data = request.get_json(force=True)
    amount = r2(data.get("amount"))
    if amount <= 0:
        return jsonify({"ok": False, "error": "amount must be positive"}), 400
    payment = Payment(receivable_id=r.id, amount=amount, note=data.get("note") or "")
    r.amount_paid = r2((r.amount_paid or 0) + amount)
    db.session.add(payment)
    db.session.commit()
    return jsonify({"ok": True, "receivable": r.to_dict()})


# ---------------------------------------------------------------------------
# Dashboard summary
# ---------------------------------------------------------------------------

@app.route("/api/dashboard")
def api_dashboard():
    by_category = {}
    total_stock_value = 0
    total_potential_profit = 0
    low_stock_count = 0
    for cat in CATEGORIES:
        items = InventoryItem.query.filter_by(category=cat, is_deleted=False).all()
        value = sum((i.quantity or 0) * (i.price_per_unit or 0) for i in items)
        profit = sum(i.total_potential_profit for i in items)
        low = sum(1 for i in items if (i.quantity or 0) <= (i.low_stock_threshold or 0))
        by_category[cat] = {
            "label": CATEGORY_LABELS[cat],
            "item_count": len(items),
            "stock_value": round(value, 2),
            "potential_profit": round(profit, 2),
            "low_stock_count": low,
        }
        total_stock_value += value
        total_potential_profit += profit
        low_stock_count += low

    receivables = Receivable.query.filter_by(is_deleted=False).all()
    total_outstanding = sum(r.balance for r in receivables if r.balance > 0)
    total_customers_owing = sum(1 for r in receivables if r.balance > 0)

    return jsonify({
        "by_category": by_category,
        "total_stock_value": round(total_stock_value, 2),
        "total_potential_profit": round(total_potential_profit, 2),
        "total_items": sum(v["item_count"] for v in by_category.values()),
        "low_stock_count": low_stock_count,
        "total_outstanding": round(total_outstanding, 2),
        "total_customers_owing": total_customers_owing,
    })


# ---------------------------------------------------------------------------
# Monthly Excel report (Inventory / Receivables / Profit - 3 sheets)
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_header(ws, headers):
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def _autosize(ws):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 40)


def build_monthly_report(year, month):
    month_start = datetime(year, month, 1)
    last_day = calendar.monthrange(year, month)[1]
    month_end = datetime(year, month, last_day, 23, 59, 59)
    month_label = month_start.strftime("%B %Y")

    wb = Workbook()

    # --- Sheet 1: Inventory (current stock, as of report generation) ---
    ws1 = wb.active
    ws1.title = "Inventory"
    ws1.append([f"Inventory Summary - as of {datetime.utcnow().strftime('%Y-%m-%d %H:%M')} UTC"])
    ws1.append([])
    _write_header(ws1, ["Category", "Item", "Quantity", "Unit", "Purchase Price",
                         "Selling Price", "Stock Value", "Profit/Unit", "Potential Profit"])
    total_stock_value = 0
    total_potential_profit = 0
    items = InventoryItem.query.filter_by(is_deleted=False).order_by(
        InventoryItem.category, InventoryItem.name
    ).all()
    for i in items:
        stock_value = round((i.quantity or 0) * (i.price_per_unit or 0), 2)
        ws1.append([
            CATEGORY_LABELS.get(i.category, i.category), i.name, i.quantity, i.unit,
            i.purchase_price, i.price_per_unit, stock_value,
            i.profit_per_unit, i.total_potential_profit,
        ])
        total_stock_value += stock_value
        total_potential_profit += i.total_potential_profit
    ws1.append([])
    ws1.append(["", "", "", "", "", "TOTAL", round(total_stock_value, 2), "", round(total_potential_profit, 2)])
    ws1[ws1.max_row][5].font = Font(bold=True)
    _autosize(ws1)

    # --- Sheet 2: Receivables (dues + payments received this month) ---
    ws2 = wb.create_sheet("Receivables")
    ws2.append([f"Receivables Summary - {month_label}"])
    ws2.append([])
    ws2.append(["Current outstanding dues (all customers, as of today)"])
    _write_header(ws2, ["Customer", "Phone", "Total Amount", "Amount Paid", "Balance", "Status", "Due Date"])
    receivables = Receivable.query.filter_by(is_deleted=False).order_by(Receivable.customer_name).all()
    total_outstanding = 0
    for r in receivables:
        ws2.append([r.customer_name, r.phone, r.total_amount, r.amount_paid, r.balance,
                    r.status, r.due_date.strftime("%Y-%m-%d") if r.due_date else ""])
        if r.balance > 0:
            total_outstanding += r.balance
    ws2.append([])
    ws2.append(["", "", "", "TOTAL OUTSTANDING", round(total_outstanding, 2)])
    ws2[ws2.max_row][3].font = Font(bold=True)

    ws2.append([])
    ws2.append([f"Payments received in {month_label}"])
    _write_header(ws2, ["Customer", "Amount Received", "Date", "Note"])
    payments = (
        Payment.query.join(Receivable)
        .filter(Payment.date >= month_start, Payment.date <= month_end)
        .order_by(Payment.date)
        .all()
    )
    total_received = 0
    for p in payments:
        ws2.append([p.receivable.customer_name if p.receivable else "-", p.amount,
                    p.date.strftime("%Y-%m-%d %H:%M") if p.date else "", p.note])
        total_received += p.amount
    ws2.append([])
    ws2.append(["", "TOTAL RECEIVED THIS MONTH", round(total_received, 2)])
    ws2[ws2.max_row][1].font = Font(bold=True)
    _autosize(ws2)

    # --- Sheet 3: Profit (snapshots logged during this month) ---
    ws3 = wb.create_sheet("Profit")
    ws3.append([f"Profit Log - {month_label}"])
    ws3.append(["Snapshots recorded whenever an item was added or updated this month."])
    ws3.append([])
    _write_header(ws3, ["Recorded At", "Category", "Item", "Quantity", "Purchase Price",
                         "Selling Price", "Profit/Unit", "Potential Profit"])
    logs = (
        ProfitLog.query.filter(ProfitLog.recorded_at >= month_start, ProfitLog.recorded_at <= month_end)
        .order_by(ProfitLog.recorded_at)
        .all()
    )
    for log in logs:
        ws3.append([
            log.recorded_at.strftime("%Y-%m-%d %H:%M") if log.recorded_at else "",
            CATEGORY_LABELS.get(log.category, log.category), log.item_name, log.quantity,
            log.purchase_price, log.price_per_unit, log.profit_per_unit, log.total_potential_profit,
        ])
    # Latest snapshot per item this month -> overall potential profit total for the month
    latest_by_item = {}
    for log in logs:
        latest_by_item[log.item_id] = log.total_potential_profit
    month_total_profit = round(sum(latest_by_item.values()), 2)
    ws3.append([])
    ws3.append(["", "", "", "", "", "TOTAL (latest snapshot per item)", "", month_total_profit])
    ws3[ws3.max_row][5].font = Font(bold=True)
    _autosize(ws3)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer, month_label


@app.route("/api/reports/monthly")
def api_monthly_report():
    month_param = request.args.get("month")  # expected "YYYY-MM"
    now = datetime.utcnow()
    try:
        if month_param:
            year, month = map(int, month_param.split("-"))
        else:
            year, month = now.year, now.month
    except (ValueError, TypeError):
        return jsonify({"ok": False, "error": "month must be in YYYY-MM format"}), 400

    buffer, month_label = build_monthly_report(year, month)
    filename = f"vendor-report-{year:04d}-{month:02d}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def ensure_db():
    with app.app_context():
        db.create_all()


# Create tables on import too (needed for WSGI/gunicorn-based hosting,
# not just "python app.py" during local development).
ensure_db()

if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
